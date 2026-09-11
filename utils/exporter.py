# utils/exporter.py
"""
Exporter 3 blok sesuai SE Dirjen Bina Konstruksi No. 47/SE/Dk/2026:

  BLOK A — AHSP per item pekerjaan (A + B + C = D, E = % × D, F = D + E)
  BLOK B — Biaya Penerapan SMKK (9 komponen A s.d I, Lampiran III)
  BLOK C — Rekapitulasi RAB/HPS (Subtotal AHSP + SMKK + PPN 11%)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import SE_NOMOR, SE_TANGGAL, DEFAULT_PPN


class Exporter:
    # ---------- Style ----------
    FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="0D47A1")
    FONT_SUB = Font(name="Calibri", size=10, italic=True, color="555555")
    FONT_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True)
    FONT_REG = Font(name="Calibri", size=10)

    FILL_HDR = PatternFill("solid", start_color="0D47A1")
    FILL_BLOCK = PatternFill("solid", start_color="BBDEFB")
    FILL_SUB = PatternFill("solid", start_color="E8F5E9")
    FILL_SMKK = PatternFill("solid", start_color="FFEBEE")
    FILL_REKAP = PatternFill("solid", start_color="FFF8E1")

    A_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    A_R = Alignment(horizontal="right", vertical="center")
    A_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

    THIN = Side(border_style="thin", color="BDBDBD")
    B_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # ======================================================
    # PUBLIC ENTRY
    # ======================================================
    @classmethod
    def export_rab_to_excel(cls, proyek_info, item_list, file_path, smkk_komponen=None):
        """
        proyek_info   : {nama_proyek, lokasi, instansi, tahun_anggaran, ppn_pct}
        item_list     : list of dict {divisi, kode_item, uraian_pekerjaan,
                                     satuan, volume, harga_dasar, overhead_pct}
        smkk_komponen : list of dict {kode, nama, satuan, koefisien, biaya}
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RAB-HPS"

        row = 1
        row = cls._tulis_header(ws, row, proyek_info)
        row = cls._tulis_blok_a(ws, row, item_list or [])
        row = cls._tulis_blok_b(ws, row, smkk_komponen or [])
        row = cls._tulis_blok_c(ws, row, item_list or [], smkk_komponen or [], proyek_info)

        widths = {1: 6, 2: 18, 3: 45, 4: 10, 5: 14, 6: 18, 7: 22}
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w

        wb.save(file_path)
        return True

    # ======================================================
    # HEADER
    # ======================================================
    @classmethod
    def _tulis_header(cls, ws, row, info):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row=row, column=1,
                value="REKAPITULASI RENCANA ANGGARAN BIAYA (RAB) / HPS PEKERJAAN KONSTRUKSI"
                ).font = cls.FONT_TITLE
        ws.cell(row=row, column=1).alignment = cls.A_C
        row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row=row, column=1,
                value=f"Berdasarkan SE Dirjen Bina Konstruksi No. {SE_NOMOR} tanggal {SE_TANGGAL}"
                ).font = cls.FONT_SUB
        ws.cell(row=row, column=1).alignment = cls.A_C
        row += 2

        for k, v in [
            ("Nama Proyek", info.get("nama_proyek", "-")),
            ("Lokasi", info.get("lokasi", "-")),
            ("Instansi", info.get("instansi", "-")),
            ("Tahun Anggaran", info.get("tahun_anggaran", "-")),
        ]:
            ws.cell(row=row, column=1, value=k).font = cls.FONT_BOLD
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
            ws.cell(row=row, column=2, value=f": {v}").font = cls.FONT_REG
            row += 1

        return row + 1

    # ======================================================
    # BLOK A — AHSP
    # ======================================================
    @classmethod
    def _tulis_blok_a(cls, ws, row, items):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1,
                    value="BLOK A — ANALISIS HARGA SATUAN PEKERJAAN (AHSP) PER ITEM")
        c.font = Font(bold=True, size=11, color="0D47A1")
        c.fill = cls.FILL_BLOCK
        c.alignment = cls.A_L
        c.border = cls.B_ALL
        row += 1

        headers = ["No", "Kode Item", "Uraian Pekerjaan", "Satuan",
                   "Volume", "Harga Satuan (F)", "Total (Vol × F)"]
        row = cls._tulis_baris_header(ws, row, headers)

        grouped = {}
        for it in items:
            grouped.setdefault(it.get("divisi", "LAIN-LAIN"), []).append(it)

        subtotal_rows = []
        no = 1

        for divisi, list_items in grouped.items():
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cd = ws.cell(row=row, column=1, value=divisi)
            cd.font = cls.FONT_BOLD
            cd.fill = cls.FILL_BLOCK
            cd.border = cls.B_ALL
            row += 1

            start = row
            for it in list_items:
                vol = float(it.get("volume", 0))
                hd = float(it.get("harga_dasar", 0))
                oh = float(it.get("overhead_pct", 10.0)) / 100.0
                harga_f = hd * (1 + oh)
                total = vol * harga_f

                cls._tulis_sel(ws, row, 1, no, cls.A_C)
                cls._tulis_sel(ws, row, 2, it.get("kode_item", ""), cls.A_C)
                cls._tulis_sel(ws, row, 3, it.get("uraian_pekerjaan", ""), cls.A_L)
                cls._tulis_sel(ws, row, 4, it.get("satuan", ""), cls.A_C)
                cls._tulis_sel(ws, row, 5, vol, cls.A_R, "#,##0.00")
                cls._tulis_sel(ws, row, 6, harga_f, cls.A_R, "#,##0.00")
                cls._tulis_sel(ws, row, 7, total, cls.A_R, "#,##0.00")
                no += 1
                row += 1
            end = row - 1

            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            ws.cell(row=row, column=1, value=f"Subtotal {divisi}").font = cls.FONT_BOLD
            ws.cell(row=row, column=1).alignment = cls.A_R
            ws.cell(row=row, column=7, value=f"=SUM(G{start}:G{end})").font = cls.FONT_BOLD
            ws.cell(row=row, column=7).number_format = "#,##0.00"
            ws.cell(row=row, column=7).alignment = cls.A_R
            for c in range(1, 8):
                ws.cell(row=row, column=c).fill = cls.FILL_SUB
                ws.cell(row=row, column=c).border = cls.B_ALL
            subtotal_rows.append(row)
            row += 2

        return row

    # ======================================================
    # BLOK B — SMKK
    # ======================================================
    @classmethod
    def _tulis_blok_b(cls, ws, row, komponen):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1,
                    value="BLOK B — BIAYA PENERAPAN SMKK (Lampiran III)")
        c.font = Font(bold=True, size=11, color="B71C1C")
        c.fill = cls.FILL_SMKK
        c.alignment = cls.A_L
        c.border = cls.B_ALL
        row += 1

        headers = ["Kode", "Komponen Biaya SMKK", "Satuan",
                   "Koefisien", "Biaya (Rp)", "", ""]
        row = cls._tulis_baris_header(ws, row, headers)

        start = row
        for k in komponen:
            cls._tulis_sel(ws, row, 1, k.get("kode", ""), cls.A_C)
            cls._tulis_sel(ws, row, 2, k.get("nama", ""), cls.A_L)
            cls._tulis_sel(ws, row, 3, k.get("satuan", "Ls"), cls.A_C)
            cls._tulis_sel(ws, row, 4, float(k.get("koefisien", 0)), cls.A_C, "0.000000")
            cls._tulis_sel(ws, row, 5, float(k.get("biaya", 0)), cls.A_R, "#,##0.00")
            ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
            row += 1
        end = row - 1 if row > start else start

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value="Subtotal Biaya SMKK (A - I)").font = cls.FONT_BOLD
        ws.cell(row=row, column=1).alignment = cls.A_R
        formula = f"=SUM(E{start}:E{end})" if end >= start else "=0"
        ws.cell(row=row, column=7, value=formula).font = cls.FONT_BOLD
        ws.cell(row=row, column=7).number_format = "#,##0.00"
        ws.cell(row=row, column=7).alignment = cls.A_R
        for c in range(1, 8):
            ws.cell(row=row, column=c).fill = cls.FILL_SMKK
            ws.cell(row=row, column=c).border = cls.B_ALL

        return row + 2

    # ======================================================
    # BLOK C — REKAP
    # ======================================================
    @classmethod
    def _tulis_blok_c(cls, ws, row, items, komponen, info):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1,
                    value="BLOK C — REKAPITULASI RAB / HPS")
        c.font = Font(bold=True, size=11, color="1B5E20")
        c.fill = cls.FILL_REKAP
        c.alignment = cls.A_L
        c.border = cls.B_ALL
        row += 1

        # Subtotal AHSP
        subtotal_ahsp = 0.0
        for it in items:
            vol = float(it.get("volume", 0))
            hd = float(it.get("harga_dasar", 0))
            oh = float(it.get("overhead_pct", 10.0)) / 100.0
            subtotal_ahsp += vol * hd * (1 + oh)

        # Subtotal SMKK
        subtotal_smkk = sum(float(k.get("biaya", 0)) for k in komponen)

        ppn_pct = float(info.get("ppn_pct", DEFAULT_PPN))
        jumlah_sebelum_ppn = subtotal_ahsp + subtotal_smkk
        ppn = jumlah_sebelum_ppn * (ppn_pct / 100.0)
        grand_total = jumlah_sebelum_ppn + ppn

        rows = [
            ("A. Subtotal AHSP (D + E)", subtotal_ahsp, False),
            ("B. Subtotal Biaya SMKK", subtotal_smkk, False),
            ("C. Jumlah Sebelum PPN (A + B)", jumlah_sebelum_ppn, True),
            (f"D. PPN ({ppn_pct:.1f}%)", ppn, False),
            ("E. GRAND TOTAL (C + D)", grand_total, True),
        ]

        for label, value, is_bold in rows:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = ws.cell(row=row, column=1, value=label)
            cell.font = cls.FONT_BOLD if is_bold else cls.FONT_REG
            cell.alignment = cls.A_R
            cval = ws.cell(row=row, column=7, value=value)
            cval.number_format = "#,##0.00"
            cval.alignment = cls.A_R
            cval.font = cls.FONT_BOLD if is_bold else cls.FONT_REG
            for c in range(1, 8):
                ws.cell(row=row, column=c).fill = cls.FILL_REKAP
                ws.cell(row=row, column=c).border = cls.B_ALL
            row += 1

        return row

    # ======================================================
    # HELPERS
    # ======================================================
    @classmethod
    def _tulis_baris_header(cls, ws, row, headers):
        for i, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=i, value=h)
            cell.font = cls.FONT_HDR
            cell.fill = cls.FILL_HDR
            cell.alignment = cls.A_C
            cell.border = cls.B_ALL
        ws.row_dimensions[row].height = 28
        return row + 1

    @classmethod
    def _tulis_sel(cls, ws, row, col, val, align, num_fmt=None):
        cell = ws.cell(row=row, column=col, value=val)
        cell.alignment = align
        cell.border = cls.B_ALL
        if num_fmt:
            cell.number_format = num_fmt